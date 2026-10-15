import csv
import io
import json
from io import BytesIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TransactionTestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from tests.fixtures import grant_view, make_orders, make_staff
from tests.models import Customer, Order
from tests.test_summary_validation import summary_payload


class SummaryExportTests(TransactionTestCase):
    def setUp(self):
        self.user = make_staff("summary-export")
        grant_view(self.user, Order, Customer)
        self.orders = make_orders(self.user)
        self.client.force_login(self.user)

    def export(self, payload, export_format="csv"):
        return self.client.post(
            reverse("django_qbe2:export"),
            {
                "definition": json.dumps(payload),
                "page_size": "1",
                "export_format": export_format,
            },
        )

    def test_malformed_summary_ordering_returns_validation_response(self):
        for ordering in (None, 1, False, [{}]):
            with self.subTest(ordering=ordering):
                response = self.export(summary_payload(ordering=ordering))
                self.assertEqual(response.status_code, 400)

    def rows(self, response):
        body = b"".join(response.streaming_content).decode()
        response.close()
        return list(csv.reader(io.StringIO(body)))

    def test_metrics_use_all_source_records_and_preserve_decimal_text(self):
        payload = summary_payload(
            metrics=[
                {"id": "m1", "function": "count", "field": None},
                {"id": "m2", "function": "sum", "field": "total"},
                {"id": "m3", "function": "avg", "field": "total"},
            ]
        )
        response = self.export(payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            self.rows(response),
            [
                ["Count records", "Sum of total", "Average of total"],
                ["2", "300.01", "150.005000"],
            ],
        )

    def test_summary_excel_preserves_boolean_and_decimal_values(self):
        payload = summary_payload(
            group_by=["active"],
            metrics=[{"id": "m1", "function": "sum", "field": "total"}],
        )
        response = self.export(payload, export_format="xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        rows = list(workbook["Report"].values)
        self.assertEqual(rows[0], ("active", "Sum of total"))
        self.assertIn((True, 300.01), rows)
        self.assertEqual(workbook["Report"]["B2"].data_type, "n")
        workbook.close()
        response.close()

    @override_settings(QBE2_MAX_ROWS=1)
    def test_summary_excel_is_bounded_and_reports_truncation(self):
        Order.objects.filter(pk=self.orders[1].pk).update(active=False)
        response = self.export(
            summary_payload(
                group_by=["active"],
                metrics=[{"id": "m1", "function": "count", "field": None}],
            ),
            export_format="xlsx",
        )
        self.assertEqual(response["X-QBE-Truncated"], "true")
        workbook = load_workbook(
            BytesIO(b"".join(response.streaming_content)), read_only=True
        )
        rows = list(workbook["Report"].values)
        self.assertEqual(rows[0], ("active", "Count records"))
        self.assertEqual(len(rows), 2)
        workbook.close()
        response.close()

    @override_settings(QBE2_MAX_ROWS=1)
    def test_group_cap_sets_header_after_complete_summary(self):
        Order.objects.filter(pk=self.orders[1].pk).update(active=False)
        response = self.export(
            summary_payload(
                group_by=["active"],
                metrics=[{"id": "m1", "function": "sum", "field": "total"}],
            )
        )
        self.assertEqual(response["X-QBE-Truncated"], "true")
        rows = self.rows(response)
        self.assertEqual(rows[0], ["active", "Sum of total"])
        self.assertEqual(len(rows), 2)
        self.assertIn(rows[1], (["false", "200.00"], ["true", "100.01"]))

    def test_groups_export_raw_codes_nulls_and_escape_formulas(self):
        first, second = self.orders
        Order.objects.filter(pk=first.pk).update(m1="A", note="=SUM(1,1)")
        Order.objects.filter(pk=second.pk).update(m1="a", note="")
        Customer.objects.filter(pk=first.customer_id).update(credit_limit=None)
        payload = summary_payload(
            group_by=["m1", "note", "customer__credit_limit"],
            metrics=[{"id": "m1", "function": "count", "field": None}],
            ordering=[{"type": "group", "field": "m1", "direction": "asc"}],
        )
        response = self.export(payload)
        rows = self.rows(response)
        self.assertEqual(rows[0], ["m1", "note", "credit limit", "Count records"])
        self.assertEqual(rows[1], ["A", "'=SUM(1,1)", "", "1"])
        self.assertEqual(rows[2], ["a", "", "", "1"])

    def test_summary_is_materialized_before_streaming(self):
        with CaptureQueriesContext(connection) as queries:
            response = self.export(summary_payload())
        self.assertGreater(len(queries), 0)
        with CaptureQueriesContext(connection) as streaming_queries:
            rows = self.rows(response)
        self.assertEqual(len(streaming_queries), 0)
        self.assertEqual(rows[1], ["2"])

    @override_settings(QBE2_MAX_AGGREGATE_ROWS=1)
    def test_execution_error_returns_before_csv_response(self):
        response = self.export(summary_payload())
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.streaming)
        self.assertContains(response, "Refine", status_code=400)

    def test_unexpected_execution_error_sends_no_csv_bytes(self):
        with patch(
            "django_qbe2.views.execute_summary",
            side_effect=ValidationError("Unavailable summary"),
        ):
            response = self.export(summary_payload())
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.streaming)
        self.assertContains(response, "Unavailable summary", status_code=400)
